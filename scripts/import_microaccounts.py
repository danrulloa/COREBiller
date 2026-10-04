"""Read-only workbook extraction into an isolated local demonstration Core.

Run with the bundled Python (openpyxl). The application itself remains dependency-free.
"""
import argparse
from collections import Counter
from datetime import date, datetime
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import sys
import warnings

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import app
import openpyxl


def clean(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def text(value):
    return '' if value is None else str(clean(value)).strip()


def import_workbook(source, destination):
    warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    book = openpyxl.load_workbook(source, data_only=True, read_only=True)
    app.initialize(destination, 'MicroCore · demostración Excel')
    counts = Counter()
    findings = []
    with app.connect(destination) as db:
        if db.execute("SELECT value FROM settings WHERE key='source_sha256'").fetchone():
            raise ValueError('Esta demostración ya fue importada. Se conserva sin sobrescribirla.')

        def archive(kind, sheet, row, reference, payload):
            db.execute('INSERT INTO source_records(kind,sheet,row_number,reference,payload) VALUES (?,?,?,?,?)',
                       (kind, sheet, row, text(reference), json.dumps(payload, ensure_ascii=False)))
            counts[kind] += 1

        tariffs = book['Tarifas']
        years = [(col, int(tariffs.cell(2,col).value)) for col in range(2, tariffs.max_column+1)
                 if isinstance(tariffs.cell(2,col).value, (int,float))]
        latest_year = max(year for _,year in years)
        latest_col = next(col for col,year in years if year==latest_year)
        numeric_rows = set()
        for row in range(3, tariffs.max_row+1):
            name = tariffs.cell(row,1).value
            if not isinstance(name,str):
                continue
            for start, year in years:
                for offset in range(4):
                    col=start+offset
                    header=tariffs.cell(1,col).value
                    # 2017 has an unlabeled Academia column. Preserve that uncertainty.
                    category=text(header) or f'Sin encabezado ({openpyxl.utils.get_column_letter(col)})'
                    value=tariffs.cell(row,col).value
                    if value is None:
                        continue
                    archive('tarifa', 'Tarifas', row, f'{year} · {category} · {name}',
                            {'Servicio':name, 'Año':year, 'Tarifa':category, 'Valor original':clean(value),
                             'Celda precio':tariffs.cell(row,col).coordinate})
                    if year==latest_year and isinstance(value,(int,float)):
                        code=f'XLS-{row:02d}-{offset+1}-{year}'
                        db.execute('INSERT INTO services(code,name,unit,category,price) VALUES (?,?,?,?,?)',
                            (code, f'{name} · {category} · {year}', 'unidad por confirmar', f'{category} · {year}', app.money(value)))
                        counts['precios_catalogo'] += 1
                        numeric_rows.add(row)
                    elif year==latest_year:
                        findings.append(f'Tarifas!{tariffs.cell(row,col).coordinate}: {name} tiene precio no numérico {value!r}; no se ofrece como tarifa calculable.')
        counts['servicios_catalogo'] = len(numeric_rows)

        clients = {}
        historical_quotes = []
        for sheet,kind in [('Copia de BDcotizaciones','cotizacion'),('BD2021','servicio')]:
            ws=book[sheet]
            rows=list(ws.iter_rows(values_only=True))
            headers=[text(v) or f'Columna {openpyxl.utils.get_column_letter(i)}' for i,v in enumerate(rows[0],1)]
            for row_number,row in enumerate(rows[1:],2):
                if not any(v is not None for v in row):
                    continue
                payload={headers[i]:clean(v) for i,v in enumerate(row)}
                archive(kind,sheet,row_number,row[2] if kind=='cotizacion' else row[1],payload)
                if kind=='cotizacion':
                    historical_quotes.append((row_number,row))
                    key=(text(row[3]),text(row[4]),text(row[9]),text(row[5]))
                    if key[0] and key[1] and key not in clients:
                        cur=db.execute('INSERT INTO clients(name,institution,email,identification) VALUES (?,?,?,?)',key)
                        clients[key]=cur.lastrowid
                        db.execute('UPDATE clients SET address=?,phone=?,city=?,tariff=? WHERE id=?',
                            (text(payload.get('Dirección')), text(payload.get('Teléfono')), text(payload.get('Ciudad')), text(payload.get('Tipo Usuario')), cur.lastrowid))
                    qty,total,price=row[18:21]
                    if all(isinstance(v,(int,float)) for v in (qty,total,price)) and abs(Decimal(str(qty))*Decimal(str(price))-Decimal(str(total)))>Decimal('.01'):
                        findings.append(f'{sheet}!fila {row_number}, {row[2]}: cantidad × valor unitario no coincide con el valor registrado. Se conserva el original sin recalcular.')
        counts['clientes'] = len(clients)
        refs=Counter(text(row[2]) for _,row in historical_quotes)
        for ref,n in refs.items():
            if n>1:
                findings.append(f'Consecutivo histórico repetido: {ref}, {n} filas. No se fusionan automáticamente.')
        for row in range(3,4):
            for col in range(2,9):
                value=book['CuerpoCorreos'].cell(row,col).value
                if value is not None:
                    archive('correo','CuerpoCorreos',row,book['CuerpoCorreos'].cell(2,col).value,
                        {'Plantilla':book['CuerpoCorreos'].cell(2,col).value,'Texto':clean(value), 'Celda':book['CuerpoCorreos'].cell(row,col).coordinate})
        for sheet in ('Descuentos','Desplegables','ListaServicios','ContadorServicios'):
            for i,row in enumerate(book[sheet].iter_rows(values_only=True),1):
                if any(v is not None for v in row):
                    archive('configuracion',sheet,i,f'{sheet} · fila {i}',
                        {openpyxl.utils.get_column_letter(j):clean(v) for j,v in enumerate(row,1) if v is not None})

        # A faithful arithmetic replay of source row 3: no current-tariff substitution.
        row_number,row=historical_quotes[1]
        key=(text(row[3]),text(row[4]),text(row[9]),text(row[5]))
        cur=db.execute('INSERT INTO services(code,name,unit,category,price) VALUES (?,?,?,?,?)',
            ('REPLAY-2021-R3',text(row[17]),'cantidad original',f'{text(row[14])} · 2021 · reproducción histórica',app.money(row[20])))
        counts['precios_reproduccion'] = 1
        metadata={'archivo':source.name,'sha256':before,'ultimo_anio_tarifas':latest_year,
                  'conteos':dict(counts),'hallazgos':findings,
                  'hojas_vacias':['BDcotizaciones','BDservicios'],
                  'nota_unidades':'El Excel no tiene una columna estructurada con unidades. El catálogo conserva unidad por confirmar.',
                  'nota_historico':'Los registros históricos se conservan como consulta de origen, no se convierten en cotizaciones emitidas ni se recalculan.'}
        for key,value in [('source_sha256',before),('source_metadata',json.dumps(metadata,ensure_ascii=False)),
                          ('terms',f'DEMOSTRACIÓN con datos de {source.name}. Tarifas de {latest_year}; no se han validado como vigentes en 2026. Unidades de cobro por confirmar. Sin validez comercial.'),
                          ('contact','MicroCore · Universidad de los Andes · demostración local')]:
            db.execute('INSERT OR REPLACE INTO settings VALUES (?,?)',(key,value))
    book.close()
    if hashlib.sha256(source.read_bytes()).hexdigest()!=before:
        raise ValueError('El archivo fuente cambió durante la importación.')
    return metadata


if __name__ == '__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--source',type=Path,default=app.ROOT/'MicroAccounts.xlsx')
    parser.add_argument('--destination',type=Path,default=app.ROOT/'data'/'microcore-excel.sqlite3')
    args=parser.parse_args()
    print(json.dumps(import_workbook(args.source,args.destination),ensure_ascii=False,indent=2))
